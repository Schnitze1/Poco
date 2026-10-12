import os
from pptx import Presentation
from pptx.util import Cm, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import win32com.client

EDA_DIR = r"d:\Repos\Poco\EDA"
ASSETS_DIR = os.path.join(EDA_DIR, "poster_assets")
POSTER_DIR = r"d:\Repos\Poco\Poster"
PPTX_OUT = os.path.join(POSTER_DIR, "POCO_A0_Poster_CL01_G01.pptx")

def build_pptx():
    print(f"Generating native A0 PowerPoint presentation: {PPTX_OUT}")
    prs = Presentation()
    prs.slide_width = Cm(118.9)
    prs.slide_height = Cm(84.1)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 1. Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(0), Cm(118.9), Cm(84.1))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg.line.fill.background()

    # 2. Row 1: Title Band (Centred)
    title_box = slide.shapes.add_textbox(Cm(1.5), Cm(1.5), Cm(115.9), Cm(5.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Poco — Preserve Originals, Conceal Only"
    r.font.name = "Calibri Light"
    r.font.size = Pt(64)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(4)
    r2 = p2.add_run()
    r2.text = "Privacy-first PII anonymisation for LLM workflows, with a fully automated multi-environment DevOps pipeline"
    r2.font.name = "Calibri Light"
    r2.font.size = Pt(28)
    r2.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(4)
    r3 = p3.add_run()
    r3.text = "SWE40006 Software Deployment and Evolution  |  Group CL01_G01  |  Tutorial: Wednesday 10:30  |  Jamie Soo (104559690) · Timothy Schabel (105334128) · Sahil (105295665)"
    r3.font.name = "Calibri Light"
    r3.font.size = Pt(20)
    r3.font.color.rgb = RGBColor(0x26, 0x26, 0x26)

    # Divider line
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.5), Cm(6.8), Cm(115.9), Cm(0.08))
    div.fill.solid()
    div.fill.fore_color.rgb = RGBColor(0x1F, 0x38, 0x64)
    div.line.fill.background()

    # 3. Row 2: Key Findings Strip
    r2_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(1.5), Cm(7.2), Cm(115.9), Cm(3.8))
    r2_bg.fill.solid()
    r2_bg.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)
    r2_bg.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)

    findings = [
        ("20 / 43", "PII labels viable for training"),
        ("8", "labels with zero validation instances"),
        ("2.4×", "more tokens per character for Chinese"),
        ("73:1", "class imbalance (GIVENNAME : TIME)"),
        ("99.99%", "of annotations in viable labels")
    ]
    bw = 115.9 / 5
    for i, (num, lbl) in enumerate(findings):
        bx = 1.5 + i * bw
        tb = slide.shapes.add_textbox(Cm(bx), Cm(7.2), Cm(bw), Cm(3.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = num + "\n"
        r.font.name = "Calibri Light"
        r.font.size = Pt(40)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        
        r2 = p.add_run()
        r2.text = lbl
        r2.font.name = "Calibri Light"
        r2.font.size = Pt(17)
        r2.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

        if i < 4:
            v_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(bx + bw), Cm(7.5), Cm(0.04), Cm(3.2))
            v_line.fill.solid()
            v_line.fill.fore_color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
            v_line.line.fill.background()

    # 4. Row 4: Demo + QR Strip
    r4_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(1.5), Cm(75.5), Cm(115.9), Cm(3.8))
    r4_bg.fill.solid()
    r4_bg.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)
    r4_bg.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)

    qw = 115.9 / 3
    q_items = [
        ("qrcode_repo.png", "Repository", "github.com/Schnitze1/SWE40006-Group-Project"),
        ("qrcode_demo.png", "Live Demo", "littlepoco.me"),
        ("qrcode_eda.png", "Full EDA Report", "Scan for full-corpus analysis")
    ]
    for qi, (q_img, q_t1, q_t2) in enumerate(q_items):
        qx = 1.5 + qi * qw + 1.0
        slide.shapes.add_picture(os.path.join(ASSETS_DIR, q_img), Cm(qx), Cm(75.8), Cm(3.2), Cm(3.2))
        tb = slide.shapes.add_textbox(Cm(qx + 3.6), Cm(75.6), Cm(qw - 4.8), Cm(3.4))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = q_t1 + "\n"
        r.font.name = "Calibri Light"
        r.font.size = Pt(22)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        r2 = p.add_run()
        r2.text = q_t2
        r2.font.name = "Calibri Light"
        r2.font.size = Pt(17)
        r2.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # 5. Row 5: Footer Strip
    foot_box = slide.shapes.add_textbox(Cm(1.5), Cm(80.0), Cm(115.9), Cm(2.2))
    tf = foot_box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Poco — Preserve Originals, Conceal Only  |  SWE40006 Software Deployment and Evolution  |  Group CL01_G01  |  October 2026"
    r.font.name = "Calibri Light"
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # Save presentation
    prs.save(PPTX_OUT)
    print("PowerPoint presentation saved successfully!")

    # Export to PDF via PowerPoint COM
    try:
        ppt_app = win32com.client.Dispatch('PowerPoint.Application')
        pres = ppt_app.Presentations.Open(os.path.abspath(PPTX_OUT), WithWindow=False)
        ppt_pdf = os.path.join(POSTER_DIR, "POCO_A0_Poster_PowerPoint.pdf")
        pres.SaveAs(os.path.abspath(ppt_pdf), 32) # 32 = ppSaveAsPDF
        pres.Close()
        ppt_app.Quit()
        print(f"Exported PowerPoint PDF successfully: {ppt_pdf}")
    except Exception as e:
        print(f"PowerPoint COM PDF export note: {e}")

if __name__ == "__main__":
    build_pptx()
